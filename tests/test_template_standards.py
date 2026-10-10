"""Template standards: ``our_standard`` is the complete clause (issue #242).

Two defects made ``vs_standard`` collapse to ``different_concept`` on a real
form: a clause split across template nodes took only its first node (a lead-in
sentence with no operative right), and a fill-in cover table classified to a
clause type became that type's standard. These tests drive both through the
pure module, the clause-type compiler and the real mine -> project pipeline.
They also pin what front matter is NOT: an operative schedule with a blank, or
any fill-in table after the form's first operative clause, stays a clause; and
front matter left out of ``our_standard`` is still the origin reference, so our
own cover-table text struck before signing is our concession. Front matter is
decided only by the real template-observation producer, over the whole
classified template (unclassified nodes included), and persisted, so every
front-matter fixture here comes from that producer and ``project`` is shown to
agree with ``mine``.

SECURITY NOTE: every fixture is synthetic text with fictional names. No real
agreement, institution or form is referenced.
"""

from __future__ import annotations

import json
from pathlib import Path

import pytest
import yaml

import playbook_engine.pipeline as pipeline
from playbook_engine.clause_classifier import ClassifiedClause, ClauseClassification
from playbook_engine.clause_position_compiler import compile_clause_positions
from playbook_engine.clause_tree import ClauseNode, ClauseTree
from playbook_engine.config import load_config
from playbook_engine.extraction import ExtractorLabel
from playbook_engine.observation_builder import (
    OUTCOME_CONCEDED_BEFORE_SIGNING,
    Observation,
    ObservationCitation,
)
from playbook_engine.pipeline import (
    _template_observations_from_classified,
    mine_corpus,
    project_playbook,
)
from playbook_engine.taxonomy import load_taxonomy
from playbook_engine.template_standards import (
    STANDARD_NODE_SEPARATOR,
    form_front_matter,
    is_form_front_matter,
    standard_observations,
    template_standards,
)
from playbook_engine.validator import validate_document

_TAXONOMY_PATH = Path(__file__).parent.parent / "spec" / "taxonomy" / "affiliation-agreement.yaml"

# --- synthetic template language ---------------------------------------------------

_COVER_TABLE = "Field | Entry\nLicensee | [Licensee Name]\nSupport tier | [Tier]\nGo-live | ____"
_PREAMBLE = (
    "Orbit Systems Ltd and the licensee identified above enter into this contract "
    "on the signature date, and Orbit Systems Ltd grants access to its platform."
)
_LEAD_IN = "9.2 Exit rights. Each side can wind up the arrangement in the following ways:"
_LIMB = (
    "9.2.1 At will. Either side can wind it up by giving the other ninety days' notice in writing."
)
_EFFECT = "9.4 Afterwards. Only the confidentiality section survives once the arrangement is over."
_MATRIX = (
    "Licensee Duties | Orbit Duties\n"
    "- Keep credentials private. | - Maintain uptime of the platform."
)

_GOVERNING_LAW = "5.1. Governing Law. Delaware law governs this Agreement."

_NEUTRAL = {"direction": "neutral", "magnitude": "none"}


def _t_obs(tid: str | None, text: str, path: str) -> Observation:
    return Observation(
        observation_id=f"template/template/{path}",
        taxonomy_id=tid,
        text_summary=text[:60],
        full_text=text,
        citation=ObservationCitation(
            document_id="template",
            version="template",
            clause_path=path,
            char_span=(int(path) * 100, int(path) * 100 + len(text)),
            version_id="template",
        ),
        deviation="none",
        risk_delta=dict(_NEUTRAL),
        provenance="our_paper",
        outcome="signed",
    )


def _cc(path: str, tid: str | None, text: str) -> ClassifiedClause:
    node = ClauseNode(clause_path=path, heading=tid, text=text, char_span=(0, max(1, len(text))))
    cls = (
        ClauseClassification(taxonomy_id=tid, confidence=0.9, basis="llm_segmenter")
        if tid
        else ClauseClassification(taxonomy_id=None, confidence=0.0, basis="unclassified")
    )
    return ClassifiedClause(node=node, classification=cls)


def _template(*specs: tuple[str, str | None, str]) -> list[Observation]:
    """Template observations from the REAL producer, over a whole classified
    template (``(clause_path, taxonomy_id or None, text)`` in document order):
    the only place front matter is decided, unclassified nodes included."""
    return _template_observations_from_classified([_cc(p, tid, text) for p, tid, text in specs])


def _deal_obs(tid: str, text: str, path: str = "1") -> Observation:
    return Observation(
        observation_id=f"deal-001/v1/{path}",
        taxonomy_id=tid,
        text_summary=text[:60],
        full_text=text,
        citation=ObservationCitation(
            document_id="deal-001", version="v1", clause_path=path, char_span=None
        ),
        deviation="substantive",
        risk_delta=dict(_NEUTRAL),
        provenance="counterparty_paper",
        outcome="signed",
    )


# ---------------------------------------------------------------------------
# is_form_front_matter
# ---------------------------------------------------------------------------


def test_fill_in_cover_table_is_front_matter() -> None:
    assert is_form_front_matter(_COVER_TABLE)


@pytest.mark.parametrize("blank", ["[Customer Name]", "____", "{{customer}}", "<<date>>"])
def test_each_placeholder_style_marks_a_fill_in_table(blank: str) -> None:
    table = f"Party | Address\n{blank} | {blank}\n{blank} | Start: signature date"
    assert is_form_front_matter(table)


def test_operative_matrix_without_blanks_is_a_clause() -> None:
    """A responsibilities matrix is a table too, but it has no fill-in fields."""
    assert not is_form_front_matter(_MATRIX)


def test_table_with_blanks_and_a_sentence_cell_is_a_clause() -> None:
    text = (
        "Fee | Terms\n"
        "[Amount] | Customer shall pay the fee within thirty days of each invoice.\n"
        "[Amount] | Late payments accrue interest."
    )
    assert not is_form_front_matter(text)


def test_prose_with_a_bracketed_blank_is_a_clause() -> None:
    assert not is_form_front_matter("The fee is [Amount] and is payable within thirty days.")


def test_markdown_link_is_not_a_blank() -> None:
    text = (
        "Party | Site\nOrbit Systems Ltd | [our terms](https://example.com/terms)\nBeta LLC | None"
    )
    assert not is_form_front_matter(text)


@pytest.mark.parametrize("text", ["", "   \n  ", _PREAMBLE, _LEAD_IN, _LIMB])
def test_ordinary_clause_text_is_not_front_matter(text: str) -> None:
    assert not is_form_front_matter(text)


def test_a_single_table_row_is_not_enough() -> None:
    assert not is_form_front_matter("[Customer Name] | [Address]")


# Operative schedules: tables whose values are mostly filled in, with ONE blank.
_COVER_SCHEDULE = (
    "Cover | Minimum limit\n"
    "Public liability | $2,000,000 each claim\n"
    "Cyber cover | [$ amount]\n"
    "Employer cover | Statutory minimum"
)
_SERVICE_LEVELS = (
    "Severity | Reply within | Fix within\n"
    "Critical | One hour | [__]\n"
    "Major | Four hours | Two days"
)


def _flatten(table: str) -> str:
    """The legacy extractor's shape of *table*: every row on ONE line."""
    return " | ".join(line.strip() for line in table.splitlines())


@pytest.mark.parametrize("table", [_COVER_SCHEDULE, _SERVICE_LEVELS])
@pytest.mark.parametrize("shape", ["rows", "flattened"])
def test_operative_schedule_with_one_blank_is_a_clause(table: str, shape: str) -> None:
    """A schedule of filled values with one blank (an insurance limit left to
    fill, an open resolution target) is an operative clause, in both shapes."""
    text = table if shape == "rows" else _flatten(table)
    assert not is_form_front_matter(text)


def test_a_fill_in_table_under_a_prose_sentence_is_a_clause() -> None:
    text = (
        "Any notice under this contract is written down and posted to the address listed here.\n"
        "Orbit Systems Ltd | [Notice address]\n"
        "Licensee | [Notice address]"
    )
    assert not is_form_front_matter(text)


# The legacy DOCX extractor joins a whole table into ONE " | "-separated line, so
# the same fill-in table reaches us as one line, not one line per row.
_FLATTENED_COVER = (
    "Field | Entry | Licensee | [Licensee Name] | Support tier | [Tier] | Go-live | ____"
)


def test_flattened_single_line_cover_table_is_front_matter() -> None:
    assert is_form_front_matter(_FLATTENED_COVER)
    assert is_form_front_matter("Order sheet\n" + _FLATTENED_COVER)


def test_flattened_matrix_without_blanks_or_with_sentences_is_a_clause() -> None:
    assert not is_form_front_matter("Licensee Duties | Orbit Duties | - Keep it safe. | - Run it.")
    assert not is_form_front_matter(
        "Fee | Terms | [Amount] | Licensee shall pay the fee within thirty days of each invoice."
    )


def test_a_table_line_among_prose_lines_is_a_clause() -> None:
    assert not is_form_front_matter(f"{_PREAMBLE}\n{_LIMB}\n{_FLATTENED_COVER}")


def _legacy_docx(tmp_path: Path, *, with_clause: bool) -> Path:
    """A synthetic DOCX: a title, a fill-in table, optionally an operative clause."""
    from docx import Document

    doc = Document()
    doc.add_paragraph("Access Order Sheet")
    rows = [
        ("Field", "Entry"),
        ("Licensee", "[Licensee Name]"),
        ("Support tier", "[Tier]"),
        ("Go-live", "____"),
    ]
    table = doc.add_table(rows=len(rows), cols=2)
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    if with_clause:
        doc.add_paragraph(
            "1. Exit. Either side can wind up the arrangement on ninety days' notice."
        )
    path = tmp_path / ("with-clause.docx" if with_clause else "cover-only.docx")
    doc.save(str(path))
    return path


def test_real_legacy_extractor_cover_table_is_front_matter(tmp_path: Path) -> None:
    """Drive a DOCX table through the real legacy extractor: it is ONE line
    (the shape docling never produces), and it must still be front matter."""
    from playbook_engine.extraction import extract_blocks

    path = _legacy_docx(tmp_path, with_clause=True)
    _, blocks, label = extract_blocks(path, extractor="legacy")
    assert str(label) == "legacy"
    table_blocks = [b.text for b in blocks if "[Licensee Name]" in b.text]
    assert len(table_blocks) == 1  # the whole table, flattened to one line
    assert "\n" not in table_blocks[0] and table_blocks[0].count("|") >= 3
    assert is_form_front_matter(table_blocks[0])
    clause_blocks = [b.text for b in blocks if b.text.startswith("1. Exit")]
    assert clause_blocks and not is_form_front_matter(clause_blocks[0])


def test_real_legacy_ingest_never_makes_the_cover_table_a_standard(tmp_path: Path) -> None:
    """The deterministic legacy path (``_ingest_file``) into template standards:
    classified to a clause type, the cover table is still a template
    observation (the origin reference) but contributes to no standard."""
    from playbook_engine.pipeline import _ingest_file

    tree = _ingest_file(_legacy_docx(tmp_path, with_clause=False), "template", "template")
    cover_nodes = [n for n in tree.nodes if "[Licensee Name]" in n.text]
    assert cover_nodes and all(is_form_front_matter(n.text) for n in cover_nodes)
    obs = _template_observations_from_classified(
        [
            _cc(n.clause_path or str(i), "parties_and_recitals", n.text)
            for i, n in enumerate(tree.nodes)
        ]
    )
    assert [o.full_text for o in form_front_matter(obs)] == [n.text for n in cover_nodes]
    assert template_standards(obs) == {}


# ---------------------------------------------------------------------------
# template_standards: split nodes are joined, front matter never wins
# ---------------------------------------------------------------------------


def test_split_clause_joins_every_node_in_document_order() -> None:
    standards = template_standards(
        [_t_obs("termination_convenience", _LEAD_IN, "7"), _t_obs("x", "x" * 30, "8")]
        + [_t_obs("termination_convenience", _LIMB, "9")]
    )
    std = standards["termination_convenience"]
    assert std.text == STANDARD_NODE_SEPARATOR.join([_LEAD_IN, _LIMB])
    assert [n.citation.clause_path for n in std.nodes] == ["7", "9"]
    assert standards["x"].text == "x" * 30


def test_single_node_standard_is_that_node_verbatim() -> None:
    std = template_standards([_t_obs("governing_law", "Delaware law governs.", "3")])
    assert std["governing_law"].text == "Delaware law governs."


def test_cover_table_never_becomes_a_standard_and_the_preamble_does() -> None:
    obs = _template(
        ("1", "parties_and_recitals", _COVER_TABLE),
        ("2", "parties_and_recitals", _PREAMBLE),
    )
    standards = template_standards(obs)
    assert standards["parties_and_recitals"].text == _PREAMBLE
    assert "Licensee Name" not in standards["parties_and_recitals"].text


def test_a_clause_type_carried_only_by_front_matter_has_no_standard() -> None:
    assert template_standards(_template(("1", "parties_and_recitals", _COVER_TABLE))) == {}


def test_the_producer_drops_unclassified_and_empty_nodes() -> None:
    """Neither an unclassified node nor an empty one becomes a template
    observation, so neither can be (part of) a standard."""
    obs = _template(("1", None, _PREAMBLE), ("2", "governing_law", "  \n"))
    assert obs == []
    assert standard_observations(obs) == []


# A fee schedule of blanks has the fill-in SHAPE; only its position decides.
_FEE_TABLE = "Service | Fee\nSet-up | [Amount]\nYearly access | [Amount]\nOn-site day | [Amount]"
_NOTICE_TABLE = (
    "Notices. Each side sends notices to its own entry here:\n"
    "Orbit Systems Ltd | [Notice address]\n"
    "Licensee | [Notice address]"
)


@pytest.mark.parametrize(
    ("tid", "table"),
    [
        ("remuneration", _FEE_TABLE),
        ("notices", _NOTICE_TABLE),
        ("remuneration", _flatten(_FEE_TABLE)),
    ],
)
def test_fill_in_table_after_the_first_operative_clause_stays_a_clause(
    tid: str, table: str
) -> None:
    assert is_form_front_matter(table)  # premise: the fill-in shape
    obs = _template(("1", "parties_and_recitals", _PREAMBLE), ("2", tid, table))
    assert form_front_matter(obs) == []
    assert template_standards(obs)[tid].text == table


def test_the_same_table_before_the_first_operative_clause_is_front_matter() -> None:
    obs = _template(
        ("1", "remuneration", _FEE_TABLE),
        ("2", "parties_and_recitals", _PREAMBLE),
    )
    assert [o.citation.clause_path for o in form_front_matter(obs)] == ["1"]
    assert "remuneration" not in template_standards(obs)


@pytest.mark.parametrize("title_tid", ["parties_and_recitals", None])
def test_a_title_node_does_not_end_the_front_matter(title_tid: str | None) -> None:
    """A title line has no sentence, classified or not: the cover table after
    it is still front matter."""
    obs = _template(
        ("1", title_tid, "Access Order Sheet"),
        ("2", "parties_and_recitals", _COVER_TABLE),
        ("3", "parties_and_recitals", _PREAMBLE),
    )
    assert [o.citation.clause_path for o in form_front_matter(obs)] == ["2"]
    assert "Licensee Name" not in template_standards(obs)["parties_and_recitals"].text


def test_an_unclassified_operative_clause_ends_the_front_matter() -> None:
    """Review repro (#242): an UNCLASSIFIED preamble sentence, then a fee table
    of blanks, then governing law. The producer still sees the unclassified
    sentence (it drops that node only after deciding front matter), so the fee
    table is an operative schedule and keeps its our_standard."""
    obs = _template(
        ("1", None, _PREAMBLE),
        ("2", "remuneration", _FEE_TABLE),
        ("3", "governing_law", _GOVERNING_LAW),
    )
    assert [o.citation.clause_path for o in obs] == ["2", "3"]  # node 1 dropped
    assert form_front_matter(obs) == []
    assert template_standards(obs)["remuneration"].text == _FEE_TABLE
    positions, _, _ = compile_clause_positions([], obs)
    std = {p.taxonomy_id: p.our_standard for p in positions}["remuneration"]
    assert std is not None and std.text == _FEE_TABLE


def test_front_matter_is_decided_over_the_whole_template_not_the_observations() -> None:
    """The same classified nodes, without the unclassified sentence ahead of
    them, make the fee table front matter: the boundary depends on a node the
    observations never carry, which is why the producer records the verdict."""
    obs = _template(("2", "remuneration", _FEE_TABLE), ("3", "governing_law", _GOVERNING_LAW))
    assert [o.citation.clause_path for o in form_front_matter(obs)] == ["2"]
    assert "remuneration" not in template_standards(obs)


def test_front_matter_flag_round_trips_through_the_store() -> None:
    """``project`` reads template observations back from
    template_observations.jsonl: the producer's verdict must survive it."""
    from playbook_engine.pipeline import _restore_observations

    obs = _template(
        ("1", "parties_and_recitals", _COVER_TABLE),
        ("2", "parties_and_recitals", _PREAMBLE),
    )
    raw = [o.to_dict() for o in obs]
    assert [r.get("x_form_front_matter") for r in raw] == [True, None]
    restored = _restore_observations(raw)
    assert [o.form_front_matter for o in restored] == [True, False]
    assert template_standards(restored) == template_standards(obs)


# ---------------------------------------------------------------------------
# compile_clause_positions: the OPF our_standard
# ---------------------------------------------------------------------------


def _our_standard(template: list[Observation], tid: str, deals: list[Observation] | None = None):
    positions, _, _ = compile_clause_positions(deals or [], template)
    return {p.taxonomy_id: p.our_standard for p in positions}[tid]


def test_our_standard_is_the_lead_in_and_its_operative_limb() -> None:
    std = _our_standard(
        [
            _t_obs("termination_convenience", _LEAD_IN, "7"),
            _t_obs("termination_convenience", _LIMB, "8"),
        ],
        "termination_convenience",
    )
    assert std is not None
    assert std.text == STANDARD_NODE_SEPARATOR.join([_LEAD_IN, _LIMB])
    assert "ninety days" in std.text


def test_joined_standard_cites_every_node_and_carries_no_span() -> None:
    std = _our_standard(
        [
            _t_obs("termination_convenience", _LEAD_IN, "7"),
            _t_obs("termination_convenience", _LIMB, "8"),
        ],
        "termination_convenience",
    )
    assert std is not None
    ref = std.source_ref.to_dict()
    assert ref == {"document_id": "template", "version": "template", "clause_path": "7, 8"}


def test_single_node_standard_keeps_its_citation_and_span() -> None:
    node = _t_obs("governing_law", "Delaware law governs.", "3")
    std = _our_standard([node], "governing_law")
    assert std is not None
    assert std.source_ref.clause_path == "3"
    assert std.source_ref.char_span == node.citation.char_span


def test_our_standard_skips_the_cover_table() -> None:
    std = _our_standard(
        _template(
            ("1", "parties_and_recitals", _COVER_TABLE),
            ("2", "parties_and_recitals", _PREAMBLE),
        ),
        "parties_and_recitals",
    )
    assert std is not None
    assert std.text == _PREAMBLE
    assert std.source_ref.clause_path == "2"


def test_front_matter_only_type_is_null_when_a_deal_evidences_it_and_absent_otherwise() -> None:
    template = _template(("1", "parties_and_recitals", _COVER_TABLE))
    positions, _, _ = compile_clause_positions([], template)
    assert positions == []  # a cover table makes no clause type appear
    deal = _deal_obs(
        "parties_and_recitals",
        "Gamma Ltd and Orbit Systems Ltd are the two parties to this contract.",
    )
    assert _our_standard(template, "parties_and_recitals", [deal]) is None


# ---------------------------------------------------------------------------
# pipeline: the template observations and the mined + projected playbook
# ---------------------------------------------------------------------------


def test_template_observations_keep_the_cover_table_and_only_the_standard_drops_it() -> None:
    """Front matter stays a template observation (the origin reference for a
    struck clause, issue #216) and leaves only the standard."""
    obs = _template_observations_from_classified(
        [
            _cc("1", "parties_and_recitals", _COVER_TABLE),
            _cc("2", "parties_and_recitals", _FLATTENED_COVER),
            _cc("3", "parties_and_recitals", _PREAMBLE),
            _cc("4", "student_responsibilities", _MATRIX),
        ]
    )
    assert [o.citation.clause_path for o in obs] == ["1", "2", "3", "4"]
    assert [o.citation.clause_path for o in form_front_matter(obs)] == ["1", "2"]
    assert template_standards(obs)["parties_and_recitals"].text == _PREAMBLE


_TEMPLATE_SPECS: list[tuple[str, str, str]] = [
    ("1", "parties_and_recitals", _COVER_TABLE),
    # The same kind of table in the shape the legacy extractor produces.
    ("2", "parties_and_recitals", _FLATTENED_COVER),
    ("3", "parties_and_recitals", _PREAMBLE),
    ("4", "termination_convenience", _LEAD_IN),
    ("5", "termination_convenience", _LIMB),
    ("6", "governing_law", _GOVERNING_LAW),
]

_DEAL_SPECS: list[tuple[str, str, str]] = [
    (
        "1",
        "parties_and_recitals",
        "Gamma LLC and Delta University contract together here; Delta University "
        "sends its learners to Gamma LLC for placements.",
    ),
    (
        "2",
        "termination_convenience",
        "Either side may wind up this contract by giving sixty days' notice in writing.",
    ),
    ("3", "governing_law", "Texas law applies to this contract and to every dispute that results."),
    ("4", "term", "The contract runs for two years starting when it is signed."),
]


def _fake_llm_segment_file(path, document_id, version, *args, **kwargs):  # noqa: ANN001, ANN002, ANN003
    specs = _TEMPLATE_SPECS if document_id == "template" else _DEAL_SPECS
    nodes = [
        ClauseNode(clause_path=p, heading=h, text=t, char_span=(0, len(t))) for p, h, t in specs
    ]
    tree = ClauseTree(
        document_id=document_id, version=version, source_file=Path(path).name, nodes=nodes
    )
    return tree, {p: h for p, h, _ in specs}, ExtractorLabel("legacy")


def test_mined_and_projected_playbook_carries_the_complete_standards(
    tmp_path: Path, monkeypatch
) -> None:
    rtf = r"{\rtf1\ansi\deff0 {\fonttbl{\f0\froman\fcharset0 Times;}}\f0\fs24 1. Term\par x.\par }"
    corpus_dir = tmp_path / "corpus"
    (corpus_dir / "deal-001").mkdir(parents=True)
    (corpus_dir / "deal-001" / "v1.rtf").write_text(rtf, encoding="utf-8")
    template_path = corpus_dir / "template.rtf"
    template_path.write_text(rtf, encoding="utf-8")
    config_path = tmp_path / "playbook.config.yaml"
    config_path.write_text(
        yaml.dump(
            {
                "agreement_type": {"id": "educational-affiliation", "name": "Affiliation"},
                "baseline": {"template": str(template_path)},
                "taxonomy": str(_TAXONOMY_PATH),
                "provenance": {"our_party_aliases": ["Orbit Systems Ltd"]},
            }
        ),
        encoding="utf-8",
    )
    out_dir = tmp_path / "out"
    cfg = load_config(config_path)
    taxonomy = load_taxonomy(_TAXONOMY_PATH)
    monkeypatch.setattr(pipeline, "_llm_segment_file", _fake_llm_segment_file)

    messages: list[str] = []
    mine_corpus(
        corpus_dir=corpus_dir,
        config=cfg,
        taxonomy=taxonomy,
        out_dir=out_dir,
        use_llm_segmentation=True,
        progress=messages.append,
    )
    project_playbook(out_dir=out_dir, config=cfg, taxonomy=taxonomy)
    # The exclusion is logged, never silent.
    assert (
        "  template standards: 6 clause(s) classified; "
        "2 form front-matter node(s) excluded from our_standard"
    ) in messages

    doc = json.loads((out_dir / "playbook.opf.json").read_text(encoding="utf-8"))
    standards = {c["taxonomy_id"]: c["our_standard"] for c in doc["evidence"]["clauses"]}
    # The cover table is not the parties clause's standard: the preamble is.
    assert standards["parties_and_recitals"]["text"] == _PREAMBLE
    # The lead-in is joined to its operative limb.
    assert standards["termination_convenience"]["text"] == STANDARD_NODE_SEPARATOR.join(
        [_LEAD_IN, _LIMB]
    )
    assert standards["termination_convenience"]["source_ref"]["clause_path"] == "4, 5"
    assert "char_span" not in standards["termination_convenience"]["source_ref"]
    assert standards["governing_law"]["text"].startswith("5.1. Governing Law.")
    assert validate_document(doc).ok
    # Both cover tables stay template observations (the origin reference);
    # neither is any clause's standard.
    template_obs = [
        json.loads(line)
        for line in (out_dir / "template_observations.jsonl").read_text("utf-8").splitlines()
    ]
    assert [o["citation"]["clause_path"] for o in template_obs] == ["1", "2", "3", "4", "5", "6"]
    # mine's front-matter verdict is persisted, so project reads the same one.
    assert [o.get("x_form_front_matter", False) for o in template_obs] == [
        True,
        True,
        False,
        False,
        False,
        False,
    ]
    assert not [s for s in standards.values() if s and "Licensee Name" in s["text"]]


# Review repro (#242) end to end: an UNCLASSIFIED preamble sentence ahead of a
# fee table of blanks. Unclassified nodes never reach template_observations.jsonl,
# so project could only agree with mine through the persisted verdict.
_UNCLASSIFIED_PREAMBLE_SPECS: list[tuple[str, str | None, str]] = [
    ("1", None, _PREAMBLE),
    ("2", "remuneration", _FEE_TABLE),
    ("3", "governing_law", _GOVERNING_LAW),
]


def _fake_unclassified_preamble_segment_file(path, document_id, version, *args, **kwargs):  # noqa: ANN001, ANN002, ANN003
    specs = _UNCLASSIFIED_PREAMBLE_SPECS if document_id == "template" else _DEAL_SPECS
    nodes = [
        ClauseNode(clause_path=p, heading=h, text=t, char_span=(0, len(t))) for p, h, t in specs
    ]
    tree = ClauseTree(
        document_id=document_id, version=version, source_file=Path(path).name, nodes=nodes
    )
    return tree, {p: h for p, h, _ in specs}, ExtractorLabel("legacy")


def test_mine_and_project_agree_when_an_unclassified_sentence_ends_the_front_matter(
    tmp_path: Path, monkeypatch
) -> None:
    rtf = r"{\rtf1\ansi\deff0 {\fonttbl{\f0\froman\fcharset0 Times;}}\f0\fs24 1. Term\par x.\par }"
    corpus_dir = tmp_path / "corpus"
    (corpus_dir / "deal-001").mkdir(parents=True)
    (corpus_dir / "deal-001" / "v1.rtf").write_text(rtf, encoding="utf-8")
    template_path = corpus_dir / "template.rtf"
    template_path.write_text(rtf, encoding="utf-8")
    config_path = tmp_path / "playbook.config.yaml"
    config_path.write_text(
        yaml.dump(
            {
                "agreement_type": {"id": "educational-affiliation", "name": "Affiliation"},
                "baseline": {"template": str(template_path)},
                "taxonomy": str(_TAXONOMY_PATH),
                "provenance": {"our_party_aliases": ["Orbit Systems Ltd"]},
            }
        ),
        encoding="utf-8",
    )
    out_dir = tmp_path / "out"
    cfg = load_config(config_path)
    taxonomy = load_taxonomy(_TAXONOMY_PATH)
    monkeypatch.setattr(pipeline, "_llm_segment_file", _fake_unclassified_preamble_segment_file)

    messages: list[str] = []
    mine_corpus(
        corpus_dir=corpus_dir,
        config=cfg,
        taxonomy=taxonomy,
        out_dir=out_dir,
        use_llm_segmentation=True,
        progress=messages.append,
    )
    assert (
        "  template standards: 2 clause(s) classified; "
        "0 form front-matter node(s) excluded from our_standard"
    ) in messages
    template_obs = [
        json.loads(line)
        for line in (out_dir / "template_observations.jsonl").read_text("utf-8").splitlines()
    ]
    assert [o["citation"]["clause_path"] for o in template_obs] == ["2", "3"]
    assert not [o for o in template_obs if o.get("x_form_front_matter")]

    doc = project_playbook(out_dir=out_dir, config=cfg, taxonomy=taxonomy)
    standards = {c["taxonomy_id"]: c["our_standard"] for c in doc["evidence"]["clauses"]}
    assert standards["remuneration"]["text"] == _FEE_TABLE
    assert standards["governing_law"]["text"] == _GOVERNING_LAW


# ---------------------------------------------------------------------------
# origin (issue #216): our struck cover-table text is our concession
# ---------------------------------------------------------------------------

# Part of the template's cover table: two of its rows, in the docling shape.
_COVER_PART = "Licensee | [Licensee Name]\nSupport tier | [Tier]"

_STRUCK_COVER_SPECS: dict[str, list[tuple[str, str, str]]] = {
    # Our first draft carries part of our cover table, the preamble and
    # governing law; the signed copy strikes the cover-table part.
    "v1": [
        ("1", "parties_and_recitals", _COVER_PART),
        ("2", "parties_and_recitals", _PREAMBLE),
        ("3", "governing_law", _GOVERNING_LAW),
    ],
    "v2": [
        ("1", "parties_and_recitals", _PREAMBLE),
        ("2", "governing_law", _GOVERNING_LAW),
    ],
}


def _fake_struck_cover_segment_file(path, document_id, version, *args, **kwargs):  # noqa: ANN001, ANN002, ANN003
    specs = _TEMPLATE_SPECS if document_id == "template" else _STRUCK_COVER_SPECS[Path(path).stem]
    nodes = [
        ClauseNode(clause_path=p, heading=h, text=t, char_span=(0, len(t))) for p, h, t in specs
    ]
    tree = ClauseTree(
        document_id=document_id, version=version, source_file=Path(path).name, nodes=nodes
    )
    return tree, {p: h for p, h, _ in specs}, ExtractorLabel("legacy")


def test_struck_cover_table_text_is_our_concession_never_a_refused_ask(
    tmp_path: Path, monkeypatch
) -> None:
    """A signed deal whose first draft carries part of our cover table, struck
    before signing: that is OUR language (front matter is left out of
    our_standard, never out of the origin reference), so the removal is
    conceded_before_signing, never proposed_then_reversed, and the deal's
    precedent record has no refused ask."""
    rtf = r"{\rtf1\ansi\deff0 {\fonttbl{\f0\froman\fcharset0 Times;}}\f0\fs24 1. Term\par x.\par }"
    corpus_dir = tmp_path / "corpus"
    deal_dir = corpus_dir / "deal-002"
    deal_dir.mkdir(parents=True)
    for name in ("v1.rtf", "v2.rtf"):
        (deal_dir / name).write_text(rtf, encoding="utf-8")
    (deal_dir / "hints.yaml").write_text("signed_version: v2.rtf\n", encoding="utf-8")
    template_path = corpus_dir / "template.rtf"
    template_path.write_text(rtf, encoding="utf-8")
    config_path = tmp_path / "playbook.config.yaml"
    config_path.write_text(
        yaml.dump(
            {
                "agreement_type": {"id": "educational-affiliation", "name": "Affiliation"},
                "baseline": {"template": str(template_path)},
                "taxonomy": str(_TAXONOMY_PATH),
                "provenance": {"our_party_aliases": ["Orbit Systems Ltd"]},
            }
        ),
        encoding="utf-8",
    )
    out_dir = tmp_path / "out"
    cfg = load_config(config_path)
    taxonomy = load_taxonomy(_TAXONOMY_PATH)
    monkeypatch.setattr(pipeline, "_llm_segment_file", _fake_struck_cover_segment_file)

    mine_corpus(
        corpus_dir=corpus_dir,
        config=cfg,
        taxonomy=taxonomy,
        out_dir=out_dir,
        use_llm_segmentation=True,
    )
    observations = [
        json.loads(line)
        for line in (out_dir / "observations.jsonl").read_text("utf-8").splitlines()
    ]
    struck = [
        o for o in observations if o["full_text"] == _COVER_PART and o["outcome"] != "opening"
    ]
    assert [o["outcome"] for o in struck] == [OUTCOME_CONCEDED_BEFORE_SIGNING]
    assert not [o for o in observations if o["outcome"] == "proposed_then_reversed"]

    doc = project_playbook(out_dir=out_dir, config=cfg, taxonomy=taxonomy)
    (record,) = [
        p for p in doc["evidence"]["precedent"] if p["taxonomy_id"] == "parties_and_recitals"
    ]
    assert record["signed"] is True
    assert record["refused_asks"] == []
    # The standard the clause is ranked against is still the preamble alone.
    (clause,) = [
        c for c in doc["evidence"]["clauses"] if c["taxonomy_id"] == "parties_and_recitals"
    ]
    assert clause["our_standard"]["text"] == _PREAMBLE
    assert clause["n_refused"] == 0
