"""L1 extraction reliability (issue #252).

Every fact in a playbook starts at L1 text extraction, and three kinds of bad
text reached it silently on the reference corpus: a DOCX that docling cannot
open (so the legacy reader ran), a failure replayed from an older docling build
forever, and scans OCR'd by the wrong recognizer (garbled, or too slow to
finish). These tests drive the real ``extract_blocks`` over real, programmatic
DOCX and PDF files and fake only the ``docling`` / ``ocrmypdf`` subprocess
seam, so what is asserted is the command the engine builds and what it does
with the result.

The docling double refuses what the real backend refuses: it raises on a DOCX
that still carries an XML comment under ``w:body`` (the real crash,
``etree.QName`` on a comment node) and records every command it was given.

SECURITY NOTE: all fixtures are built at test runtime (python-docx, fpdf2,
Pillow). Fictional names only ("Alpha Corp", "Beta Ltd").
"""

from __future__ import annotations

import json
import subprocess
import zipfile
from pathlib import Path
from typing import Any

import pytest
import yaml
from docx import Document
from lxml import etree

from playbook_engine import extraction
from playbook_engine.config import load_config
from playbook_engine.extraction import ExtractionCache, ExtractionError, extract_blocks
from playbook_engine.inspection_report import build_inspection_report
from playbook_engine.pipeline import mine_corpus
from playbook_engine.taxonomy import load_taxonomy

try:
    from fpdf import FPDF  # type: ignore[import-untyped]
    from PIL import Image

    _PDF_TOOLS = True
except ImportError:
    _PDF_TOOLS = False

needs_pdf_tools = pytest.mark.skipif(not _PDF_TOOLS, reason="fpdf2/Pillow not installed")

_TAXONOMY_PATH = Path(__file__).parent.parent / "spec" / "taxonomy" / "nda.yaml"


# ---------------------------------------------------------------------------
# The docling / ocrmypdf double
# ---------------------------------------------------------------------------


class _Docling:
    """Replaces ``subprocess.run`` for docling (and ocrmypdf) and ``shutil.which``.

    ``calls`` holds every command; ``timeouts`` the ``timeout=`` each docling
    call was given; ``targets`` the file each docling call was pointed at.
    """

    def __init__(self, monkeypatch: pytest.MonkeyPatch, *, markdown: str = "# Heading\n\nBody.\n"):
        self.markdown = markdown
        self.calls: list[list[str]] = []
        self.timeouts: list[float] = []
        self.targets: list[Path] = []
        self.stderr: str | None = None  # when set, docling fails with this stderr
        self.ocrmypdf_text: list[str] | None = None  # when set, ocrmypdf "recovers" it
        monkeypatch.setattr(
            extraction.shutil,
            "which",
            lambda cmd: f"/usr/bin/{cmd}" if cmd in ("docling", "ocrmypdf") else None,
        )
        monkeypatch.setattr(extraction.subprocess, "run", self._run)

    def _run(self, cmd: list[str], **kwargs: Any) -> subprocess.CompletedProcess[str]:
        self.calls.append(cmd)
        if cmd[0] == "ocrmypdf":
            if self.ocrmypdf_text is None:
                raise subprocess.CalledProcessError(1, cmd, stderr="no ocr")
            _write_text_pdf(Path(cmd[-1]), self.ocrmypdf_text)
            return subprocess.CompletedProcess(cmd, 0, stdout="", stderr="")
        assert cmd[0] == "docling", cmd
        assert kwargs.get("timeout"), "docling must always run under a timeout"
        self.timeouts.append(kwargs["timeout"])
        target = Path(cmd[2])
        self.targets.append(target)
        if self.stderr is not None:
            raise subprocess.CalledProcessError(1, cmd, stderr=self.stderr)
        if target.suffix == ".docx":
            with zipfile.ZipFile(target) as z:
                if b"<!--" in z.read("word/document.xml"):
                    # What docling's Word backend really does.
                    raise subprocess.CalledProcessError(
                        1,
                        cmd,
                        stderr="ValueError: Invalid input tag of type cython_function_or_method",
                    )
        outdir = Path(cmd[cmd.index("--output") + 1])
        (outdir / f"{target.stem}.md").write_text(self.markdown, encoding="utf-8")
        return subprocess.CompletedProcess(cmd, 0, stdout="", stderr="")

    @property
    def docling_calls(self) -> list[list[str]]:
        return [c for c in self.calls if c[0] == "docling"]


# ---------------------------------------------------------------------------
# Fixtures: real DOCX and PDF files
# ---------------------------------------------------------------------------


def _docx(tmp_path: Path, name: str = "doc.docx", *, xml_comment: bool = False) -> Path:
    doc = Document()
    doc.add_heading("Confidentiality", level=1)
    doc.add_paragraph("Alpha Corp shall keep Beta Ltd information secret.")
    if xml_comment:
        doc.element.body.insert(1, etree.Comment(" generated by the form template "))
    path = tmp_path / name
    doc.save(str(path))
    return path


def _write_text_pdf(dest: Path, lines: list[str]) -> None:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    for line in lines:
        pdf.cell(0, 10, line, new_x="LMARGIN", new_y="NEXT")
    pdf.output(str(dest))


def _digital_pdf(tmp_path: Path, name: str, *, pages: int, words_per_page: int) -> Path:
    """A born-digital PDF: a real text layer and no images."""
    pdf = FPDF()
    pdf.set_font("Helvetica", size=8)
    for _ in range(pages):
        pdf.add_page()
        if words_per_page:
            pdf.multi_cell(0, 5, " ".join(["clause"] * words_per_page))
    dest = tmp_path / name
    pdf.output(str(dest))
    return dest


def _scan_pdf(tmp_path: Path, name: str, *, pages: int, stamp_words: int = 0) -> Path:
    """A scan: every page is one full-page image. ``stamp_words`` > 0 adds a
    small text layer on each page, as a signature stamp or scanner header does."""
    image = tmp_path / "page.png"
    Image.new("RGB", (300, 400), "white").save(image)
    pdf = FPDF()
    pdf.set_font("Helvetica", size=8)
    for _ in range(pages):
        pdf.add_page()
        pdf.image(str(image), x=0, y=0, w=pdf.w, h=pdf.h)
        if stamp_words:
            pdf.set_xy(10, 10)
            pdf.multi_cell(0, 4, " ".join(["stamp"] * stamp_words))
    dest = tmp_path / name
    pdf.output(str(dest))
    return dest


# ---------------------------------------------------------------------------
# Class A: a DOCX with an XML comment under w:body
# ---------------------------------------------------------------------------


def test_docx_with_a_body_level_xml_comment_is_extracted_by_docling(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The double crashes on any DOCX that still has the comment, as docling
    does. The engine must therefore hand docling a copy without it, and the
    version must come out ``docling`` with no fallback reason, not ``legacy``."""
    docling = _Docling(monkeypatch, markdown="# Confidentiality\n\nAlpha Corp keeps it secret.\n")
    path = _docx(tmp_path, xml_comment=True)
    original = path.read_bytes()

    _text, _blocks, label = extract_blocks(path)

    assert label == "docling"
    assert label.reason is None, "no legacy fallback may be needed for this file"
    assert len(docling.docling_calls) == 1, "one docling run, no retry"
    target = docling.targets[0]
    assert target != path and target.name == path.name, "converted from a staged copy"
    assert path.read_bytes() == original, "the corpus file itself is never modified"


def test_docx_without_xml_comments_is_converted_in_place(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The variant that must NOT change behaviour: a plain DOCX is handed to
    docling as-is, so the common case pays nothing."""
    docling = _Docling(monkeypatch)
    path = _docx(tmp_path)

    _text, _blocks, label = extract_blocks(path)

    assert label == "docling"
    assert docling.targets == [path]


def test_xml_comment_docx_under_the_double_really_does_crash_without_the_prepass(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Guards the test above against a vacuous pass: with the pre-pass disabled
    the same file takes the old route (crash, normalized retry, legacy)."""
    monkeypatch.setattr(extraction, "strip_xml_comments_docx", lambda path, dest=None: None)
    path = _docx(tmp_path, xml_comment=True)

    def _copy_without_stripping(source: Path) -> Path:
        # The retry's temp copy (the caller deletes it), as the normalizer
        # produced it before it learned to drop XML comments.
        copy = tmp_path / "retry-copy.docx"
        copy.write_bytes(source.read_bytes())
        return copy

    monkeypatch.setattr(extraction, "normalize_tracked_docx", _copy_without_stripping)
    _Docling(monkeypatch)

    _text, _blocks, label = extract_blocks(path)

    assert label == "legacy"
    assert label.reason == "backend-error"


# ---------------------------------------------------------------------------
# Class B: the cache must know which docling stack wrote an entry
# ---------------------------------------------------------------------------


def test_extractor_stack_names_the_docling_release_and_the_ocr_setup(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(extraction, "_docling_version", lambda: "2.140.0")
    stack = extraction.extractor_stack("docling")
    assert "2.140.0" in stack
    assert "tesseract" in stack
    assert "iso:en" in stack
    assert extraction.extractor_stack("legacy") == ""
    monkeypatch.setattr(extraction, "_docling_version", lambda: "2.113.0")
    assert extraction.extractor_stack("docling") != stack


def test_a_backend_error_from_an_older_docling_is_not_replayed(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Class B. docling 2.113 crashed on this file, so its fallback was cached
    as ``backend-error``. A newer docling converts it fine; the old failure must
    not be served again."""
    docling = _Docling(monkeypatch)
    docling.stderr = "boom from the old build"
    path = _docx(tmp_path)
    cache = ExtractionCache(tmp_path / "extraction_cache.jsonl")
    monkeypatch.setattr(extraction, "_docling_version", lambda: "2.113.0")

    _t, _b, old = extract_blocks(path, cache=cache)
    assert (old, old.reason) == ("legacy", "backend-error")

    # Same stack: replayed from the cache, docling not asked again.
    n = len(docling.docling_calls)
    _t, _b, replay = extract_blocks(path, cache=cache)
    assert (replay, replay.reason) == ("legacy", "backend-error")
    assert len(docling.docling_calls) == n

    # Newer docling: the entry is not a hit; docling runs and now succeeds.
    docling.stderr = None
    monkeypatch.setattr(extraction, "_docling_version", lambda: "2.140.0")
    _t, _b, fresh = extract_blocks(path, cache=cache)
    assert (fresh, fresh.reason) == ("docling", None)
    assert len(docling.docling_calls) == n + 1

    # And that result is cached under the new stack.
    _t, _b, again = extract_blocks(path, cache=cache)
    assert again == "docling"
    assert len(docling.docling_calls) == n + 1


@needs_pdf_tools
def test_a_cached_no_text_failure_is_scoped_to_its_docling_stack(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A negative-cached failure is replayed under the same stack and retried
    under another; and it names the docling failure instead of an empty page."""
    docling = _Docling(monkeypatch, markdown="<!-- image -->\n")  # docling reads nothing
    path = _scan_pdf(tmp_path, "scan.pdf", pages=1)
    cache = ExtractionCache(tmp_path / "extraction_cache.jsonl")
    monkeypatch.setattr(extraction, "_docling_version", lambda: "2.113.0")

    with pytest.raises(ExtractionError):
        extract_blocks(path, cache=cache)
    assert cache.get_failure(path, extractor="docling") is not None
    n = len(docling.docling_calls)

    with pytest.raises(ExtractionError, match="cached failure"):
        extract_blocks(path, cache=cache)
    assert len(docling.docling_calls) == n, "same stack: the failure is replayed"

    monkeypatch.setattr(extraction, "_docling_version", lambda: "2.140.0")
    assert cache.get_failure(path, extractor="docling") is None
    with pytest.raises(ExtractionError):
        extract_blocks(path, cache=cache)
    assert len(docling.docling_calls) == n + 1, "another stack: docling runs again"


def test_an_entry_with_no_stack_is_not_a_docling_hit(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The shape every cache written before this change holds: docling
    environment, no ``extractor_stack``. Nothing says which docling (or which
    recognizer) made it, so it is not trusted."""
    docling = _Docling(monkeypatch)
    path = _docx(tmp_path)
    cache = ExtractionCache(tmp_path / "extraction_cache.jsonl")
    text = "Garbled text from an older recognizer."
    cache._store.put(
        extraction._extraction_cache_payload(path, "docling"),
        {
            "canonical_text": text,
            "blocks": [{"block_id": "b0", "page": 0, "char_span": [0, len(text)]}],
            "extractor": "docling",
            "reason": None,
            "fallback_from": None,
        },
    )
    assert cache.get(path, extractor="docling") is None

    canonical, _blocks, label = extract_blocks(path, cache=cache)

    assert label == "docling"
    assert canonical != text
    assert len(docling.docling_calls) == 1


def _plant_stale_format_docling_entry(
    cache: ExtractionCache, path: Path, text: str, **extra: object
) -> None:
    """A docling-environment success filed under format 4, which the ladder can
    carry forward, with *extra* fields (``extractor_stack``) as given."""
    cache._store.put(
        {
            "file_sha256": extraction._sha256_file(path),
            "format_version": "4",
            "extractor_env": "docling",
        },
        {
            "canonical_text": text,
            "blocks": [{"block_id": "b0", "page": 0, "char_span": [0, len(text)]}],
            "extractor": "docling",
            "reason": None,
            "fallback_from": None,
            **extra,
        },
    )


def test_a_stale_format_entry_with_no_stack_is_invalidated_not_migrated(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The format ladder would carry this entry forward, but the stack check
    refuses it, so it is re-extracted. The counters must say so: reporting it
    as migrated "with no re-extraction" (what ``playbook segment`` prints) would
    be false for a file that is extracted again."""
    docling = _Docling(monkeypatch)
    path = _docx(tmp_path)
    cache = ExtractionCache(tmp_path / "extraction_cache.jsonl")
    _plant_stale_format_docling_entry(cache, path, "Garbled text from an older recognizer.")

    assert cache.get(path, extractor="docling") is None
    assert cache.migrated_count == 0
    assert cache.invalidated_count == 1

    # A second lookup (extract_blocks asks twice) neither re-counts nor re-files it.
    assert cache.get_failure(path, extractor="docling") is None
    assert (cache.migrated_count, cache.invalidated_count) == (0, 1)

    canonical, _blocks, label = extract_blocks(path, cache=cache)
    assert label == "docling"
    assert canonical != "Garbled text from an older recognizer."
    assert len(docling.docling_calls) == 1


def test_a_stale_format_entry_with_the_current_stack_is_still_migrated(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The other side of the check above: the same entry carrying the current
    stack is carried forward without re-extraction."""
    docling = _Docling(monkeypatch)
    path = _docx(tmp_path)
    cache = ExtractionCache(tmp_path / "extraction_cache.jsonl")
    text = "Alpha Corp shall indemnify Beta Ltd."
    _plant_stale_format_docling_entry(
        cache, path, text, extractor_stack=extraction.extractor_stack("docling")
    )

    hit = cache.get(path, extractor="docling")

    assert hit is not None
    assert hit[0] == text
    assert (cache.migrated_count, cache.invalidated_count) == (1, 0)
    assert docling.docling_calls == []


def test_a_legacy_environment_entry_does_not_depend_on_the_docling_stack(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The legacy reader has no docling, so a docling upgrade leaves its cache
    alone (the other variant of the branch above)."""
    monkeypatch.setattr(extraction.shutil, "which", lambda cmd: None)
    path = _docx(tmp_path)
    cache = ExtractionCache(tmp_path / "extraction_cache.jsonl")
    monkeypatch.setattr(extraction, "_docling_version", lambda: "2.113.0")
    extract_blocks(path, cache=cache, extractor="legacy")
    monkeypatch.setattr(extraction, "_docling_version", lambda: "2.140.0")
    assert cache.get(path, extractor="legacy", declared_extractor="legacy") is not None


def test_a_cached_fallback_is_replayed_not_retried(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A file docling cannot convert falls back once; the fallback is cached
    and replayed (with its reason) instead of re-running docling every time."""
    docling = _Docling(monkeypatch)
    docling.stderr = "docling: conversion failed"
    path = _docx(tmp_path)
    cache = ExtractionCache(tmp_path / "extraction_cache.jsonl")

    _t, _b, first = extract_blocks(path, cache=cache)
    n = len(docling.docling_calls)
    _t, _b, second = extract_blocks(path, cache=cache)

    assert (first, first.reason) == ("legacy", "backend-error")
    assert (second, second.reason, second.fallback_from) == ("legacy", "backend-error", "docling")
    assert len(docling.docling_calls) == n


# ---------------------------------------------------------------------------
# OCR engine and language reach docling; an unavailable engine stays loud
# ---------------------------------------------------------------------------


@needs_pdf_tools
def test_ocr_engine_and_language_reach_docling_in_the_current_api_form(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    docling = _Docling(monkeypatch)
    extract_blocks(_digital_pdf(tmp_path, "d.pdf", pages=1, words_per_page=120))
    cmd = docling.docling_calls[0]
    assert cmd[cmd.index("--ocr-engine") + 1] == "tesseract"
    assert cmd[cmd.index("--ocr-lang") + 1] == "iso:en"
    assert "eng" not in cmd, "the bare language code current docling rejects"


@needs_pdf_tools
def test_an_unavailable_ocr_engine_is_recorded_not_swallowed(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """If docling cannot start its OCR engine the version must say so. A
    digital PDF still reads through the legacy path, labelled as a fallback
    with the docling error; a scan with nothing to fall back to raises with the
    docling error in the message rather than passing for a blank page."""
    docling = _Docling(monkeypatch)
    docling.stderr = "OCR engine 'tesseract' is not installed"

    digital = _digital_pdf(tmp_path, "d.pdf", pages=1, words_per_page=120)
    _t, _b, label = extract_blocks(digital)
    assert (label, label.reason, label.fallback_from) == ("legacy", "backend-error", "docling")
    assert label.detail is not None and "tesseract" in label.detail

    scan = _scan_pdf(tmp_path, "scan.pdf", pages=1)
    with pytest.raises(ExtractionError, match="docling failed.*tesseract") as excinfo:
        extract_blocks(scan)
    assert excinfo.value.reason == extraction.FAILURE_NO_TEXT


# ---------------------------------------------------------------------------
# Classes C and D: thin text layers get full-page OCR; timeouts scale with pages
# ---------------------------------------------------------------------------


@needs_pdf_tools
@pytest.mark.parametrize(
    ("label", "build", "forced"),
    [
        # A stamp over scanned pages: the case docling skipped OCR for.
        (
            "scan_with_thin_stamp_layer",
            lambda p: _scan_pdf(p, "a.pdf", pages=4, stamp_words=5),
            True,
        ),
        (
            "one_page_scan_with_a_header",
            lambda p: _scan_pdf(p, "b.pdf", pages=1, stamp_words=8),
            True,
        ),
        # No text at all: a plain scan, docling's default OCR already reads it.
        ("plain_scan_without_text", lambda p: _scan_pdf(p, "c.pdf", pages=3), False),
        # Scanned pages whose text layer is complete: nothing to replace.
        (
            "scan_with_full_text_layer",
            lambda p: _scan_pdf(p, "d.pdf", pages=2, stamp_words=300),
            False,
        ),
        # A short born-digital page must not be re-read through OCR.
        (
            "short_digital_page",
            lambda p: _digital_pdf(p, "e.pdf", pages=1, words_per_page=30),
            False,
        ),
        (
            "long_digital_document",
            lambda p: _digital_pdf(p, "f.pdf", pages=3, words_per_page=200),
            False,
        ),
    ],
)
def test_a_thin_text_layer_over_scanned_pages_gets_full_page_ocr(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    label: str,
    build: Any,
    forced: bool,
) -> None:
    docling = _Docling(monkeypatch)
    path = build(tmp_path)
    profile = extraction._profile_pdf(path)
    assert profile is not None, label

    extract_blocks(path)

    cmd = docling.docling_calls[0]
    assert profile.force_full_page_ocr is forced, label
    assert ("--ocr-mode" in cmd) is forced, label
    if forced:
        assert cmd[cmd.index("--ocr-mode") + 1] == "full_page"


@needs_pdf_tools
def test_an_unreadable_pdf_gets_no_profile_and_the_flat_timeout(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    docling = _Docling(monkeypatch)
    broken = tmp_path / "broken.pdf"
    broken.write_bytes(b"%PDF-1.4 not really a pdf")
    assert extraction._profile_pdf(broken) is None

    extract_blocks(broken)

    assert docling.timeouts == [extraction._DOCLING_TIMEOUT_S]
    assert "--ocr-mode" not in docling.docling_calls[0]


@needs_pdf_tools
def test_the_docling_timeout_grows_with_the_page_count(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    docling = _Docling(monkeypatch)
    ten = _scan_pdf(tmp_path, "ten.pdf", pages=10)
    two = _scan_pdf(tmp_path, "two.pdf", pages=2)

    extract_blocks(ten)
    extract_blocks(two)
    docx = _docx(tmp_path)
    extract_blocks(docx)

    base, per_page = extraction._DOCLING_PDF_BASE_S, extraction._DOCLING_PDF_PER_PAGE_S
    assert docling.timeouts == [
        base + per_page * 10,  # 720 s for the 10-page scan that timed out at 600 s
        extraction._DOCLING_TIMEOUT_S,  # a short PDF never gets less than the flat cap
        extraction._DOCLING_TIMEOUT_S,  # and a DOCX keeps it
    ]
    assert docling.timeouts[0] > extraction._DOCLING_TIMEOUT_S


def test_the_timeout_has_a_ceiling() -> None:
    profile = extraction._PdfProfile
    assert extraction._docling_timeout_s(profile(pages=30, force_full_page_ocr=False)) == 1920
    assert (
        extraction._docling_timeout_s(profile(pages=5000, force_full_page_ocr=False))
        == extraction._DOCLING_PDF_CEILING_S
    )
    assert extraction._docling_timeout_s(None) == extraction._DOCLING_TIMEOUT_S


@needs_pdf_tools
def test_a_timeout_stays_recorded_and_is_never_cached(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A timeout is a fact about the run: the label says it and the cache does
    not keep the fallback, so the next run (with the page-scaled cap) retries."""
    docling = _Docling(monkeypatch)
    path = _digital_pdf(tmp_path, "d.pdf", pages=1, words_per_page=120)
    cache = ExtractionCache(tmp_path / "extraction_cache.jsonl")

    def _timeout(cmd: list[str], **kwargs: Any) -> subprocess.CompletedProcess[str]:
        raise subprocess.TimeoutExpired(cmd, kwargs["timeout"])

    monkeypatch.setattr(extraction.subprocess, "run", _timeout)
    _t, _b, label = extract_blocks(path, cache=cache)
    assert label.timed_out is True
    assert cache.get(path, extractor="docling") is None

    monkeypatch.setattr(extraction.subprocess, "run", docling._run)
    _t, _b, retry = extract_blocks(path, cache=cache)
    assert (retry, retry.timed_out) == ("docling", False)


@needs_pdf_tools
def test_ocrmypdf_second_path_may_open_a_signed_scan(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """ocrmypdf refuses a digitally signed PDF unless told it may invalidate the
    signature on the temp copy it reads (issue #252)."""
    docling = _Docling(monkeypatch, markdown="<!-- image -->\n")
    docling.ocrmypdf_text = ["Alpha Corp shall keep Beta Ltd information secret."]
    path = _scan_pdf(tmp_path, "scan.pdf", pages=1)

    _t, _b, label = extract_blocks(path)

    assert label.reason == "ocr-recovered"
    ocr_cmd = next(c for c in docling.calls if c[0] == "ocrmypdf")
    assert "--invalidate-digital-signatures" in ocr_cmd


# ---------------------------------------------------------------------------
# The thin-extraction guard
# ---------------------------------------------------------------------------


def _row(version: str, words: Any, status: str = "ok") -> dict[str, Any]:
    return {"version": version, "status": status, "words": words}


def test_thin_extraction_versions_flags_a_version_under_a_fifth_of_the_median() -> None:
    rows = [_row("v1", 5100), _row("v2", 5000), _row("v3", 61)]
    assert extraction.thin_extraction_versions(rows) == ["v3"]


def test_thin_extraction_versions_boundaries() -> None:
    thin = extraction.thin_extraction_versions
    # Exactly a fifth of the median is not flagged; just under is.
    assert thin([_row("a", 1000), _row("b", 1000), _row("c", 200)]) == []
    assert thin([_row("a", 1000), _row("b", 1000), _row("c", 199)]) == ["c"]
    # One counted version has nothing to be compared with.
    assert thin([_row("a", 3)]) == []
    # A failed version and a row without a count are ignored, not read as zero.
    assert thin([_row("a", 1000), _row("b", 1000), _row("c", None), _row("d", 0, "failed")]) == []
    # A median of zero flags nothing.
    assert thin([_row("a", 0), _row("b", 0), _row("c", 0)]) == []
    # A short sibling does not hide itself: the median, not the mean.
    assert thin([_row("a", 4000), _row("b", 4200), _row("c", 100), _row("d", 90)]) == ["c", "d"]


_RTF_PROLOGUE = r"{\rtf1\ansi\deff0{\fonttbl{\f0\froman\fcharset0 Times New Roman;}}\f0\fs24 "
_FULL = (
    r"1. Confidentiality\par Each party shall keep the other party's confidential "
    r"information secret for {n} years and shall not disclose it to any third party "
    r"without prior written consent.\par 2. Governing Law\par This Agreement is governed "
    r"by the laws of the State of California and the parties submit to its courts.\par "
)
_SHORT = r"1. Confidentiality\par Keep it secret.\par "


def _mine_thin_deal(tmp_path: Path, third: str = _SHORT) -> tuple[Path, list[str]]:
    corpus = tmp_path / "corpus"
    deal = corpus / "deal-x"
    deal.mkdir(parents=True, exist_ok=True)
    for stem, body in (
        ("v1", _FULL.replace("{n}", "three")),
        ("v2", _FULL.replace("{n}", "five")),
        ("v3", third),
    ):
        (deal / f"{stem}.rtf").write_text(_RTF_PROLOGUE + body + "}", encoding="utf-8")
    cfg = {
        "agreement_type": {"id": "nda", "name": "Mutual Non-Disclosure Agreement"},
        "baseline": {"template": None},
        "taxonomy": str(_TAXONOMY_PATH),
        "provenance": {"our_party_aliases": ["Alpha Corp"]},
    }
    config_path = tmp_path / "playbook.config.yaml"
    config_path.write_text(yaml.dump(cfg), encoding="utf-8")
    out_dir = tmp_path / "out"
    lines: list[str] = []
    mine_corpus(
        corpus_dir=corpus,
        config=load_config(config_path),
        taxonomy=load_taxonomy(_TAXONOMY_PATH),
        out_dir=out_dir,
        progress=lines.append,
    )
    return out_dir, lines


def test_mine_records_word_counts_and_flags_the_thin_version(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Through the real pipeline: the manifest rows carry ``words`` as the
    producer writes them, the run warns about the thin version, and the
    inspection report lists it."""
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    out_dir, lines = _mine_thin_deal(tmp_path)

    manifest = json.loads((out_dir / "corpus_manifest.json").read_text(encoding="utf-8"))
    rows = {r["version"]: r for r in manifest[0]["version_ingest"]}
    assert all(isinstance(r["words"], int) and r["words"] > 0 for r in rows.values())
    assert rows["v3"]["words"] < 0.2 * rows["v1"]["words"]
    assert all(r["docling_timed_out"] is False for r in rows.values())

    assert extraction.thin_extraction_versions(manifest[0]["version_ingest"]) == ["v3"]
    assert any("WARNING" in ln and "v3" in ln and "fewer words" in ln for ln in lines)
    report = build_inspection_report(out_dir)
    assert "version_ingest_thin" in report


def test_a_deal_with_no_thin_version_is_not_flagged(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    out_dir, lines = _mine_thin_deal(tmp_path, third=_FULL.replace("{n}", "one"))
    manifest = json.loads((out_dir / "corpus_manifest.json").read_text(encoding="utf-8"))
    assert extraction.thin_extraction_versions(manifest[0]["version_ingest"]) == []
    assert not any("fewer words" in ln for ln in lines)
    assert "version_ingest_thin" not in build_inspection_report(out_dir)


def test_a_docling_upgrade_busts_the_per_version_stage_cache(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """L1 results are cached per version file. Upgrading docling (or changing its
    OCR setup) changes the text for identical bytes, so every version must get a
    new L1 key; an unchanged stack must reuse them."""
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    monkeypatch.setattr(
        extraction.shutil, "which", lambda cmd: "/usr/bin/docling" if cmd == "docling" else None
    )
    index = tmp_path / "out" / ".cache" / "index.json"

    def l1_keys() -> set[str]:
        return set(json.loads(index.read_text(encoding="utf-8")).keys())

    monkeypatch.setattr(extraction, "_docling_version", lambda: "2.113.0")
    _mine_thin_deal(tmp_path)
    old = l1_keys()

    _mine_thin_deal(tmp_path)  # same stack: pure replay
    assert l1_keys() == old

    monkeypatch.setattr(extraction, "_docling_version", lambda: "2.140.0")
    _mine_thin_deal(tmp_path)
    assert len(l1_keys() - old) >= 3, "one new L1 key per version after a docling upgrade"
